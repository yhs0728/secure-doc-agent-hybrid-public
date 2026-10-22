import csv
import io
import os
import zipfile

from docx import Document
from pypdf import PdfReader

SUPPORTED = {".txt", ".pdf", ".docx", ".csv"}
# Python's default 128 KiB cell limit is too small for document-text CSVs.
CSV_FIELD_LIMIT = 8_000_000
csv.field_size_limit(CSV_FIELD_LIMIT)


def parse_and_split(content, extension):
    if extension == ".pdf":
        pdf = PdfReader(io.BytesIO(content))
        pages = ((i + 1, page.extract_text() or "") for i, page in enumerate(pdf.pages))
    elif extension == ".docx":
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if (
                len(entries) > 10000
                or sum(info.file_size for info in entries) > 100 * 1024 * 1024
            ):
                raise ValueError("DOCX_EXPANSION_LIMIT_EXCEEDED")
        doc = Document(io.BytesIO(content))
        lines = [p.text for p in doc.paragraphs]
        lines.extend(
            " | ".join(cell.text for cell in row.cells)
            for table in doc.tables
            for row in table.rows
        )
        pages = [(1, "\n".join(lines))]
    elif extension == ".csv":
        rows = csv.reader(io.StringIO(content.decode("utf-8-sig")))
        try:
            pages = [(1, "\n".join(" | ".join(row) for row in rows))]
        except csv.Error as exc:
            code = (
                "CSV_FIELD_LIMIT_EXCEEDED"
                if "field limit" in str(exc)
                else "INVALID_CSV"
            )
            raise ValueError(code) from exc
    elif extension == ".txt":
        pages = [(1, content.decode("utf-8-sig"))]
    else:
        raise ValueError("UNSUPPORTED_FILE_TYPE")
    size = int(os.getenv("CHUNK_SIZE", "500"))
    overlap = int(os.getenv("CHUNK_OVERLAP", "50"))
    if size < 1 or not 0 <= overlap < size:
        raise ValueError("INVALID_CHUNK_CONFIGURATION")
    max_chars = int(os.getenv("MAX_EXTRACTED_CHARACTERS", "8000000"))
    max_chunks = int(os.getenv("MAX_DOCUMENT_CHUNKS", "10000"))
    if max_chars < 1 or max_chunks < 1:
        raise ValueError("INVALID_DOCUMENT_LIMITS")
    chunks = []
    characters = 0
    page_count = 0
    for page, text in pages:
        page_count = page
        characters += len(text)
        if characters > max_chars:
            raise ValueError("EXTRACTED_TEXT_LIMIT_EXCEEDED")
        text = text.strip()
        for start in range(0, len(text), size - overlap):
            chunk = text[start : start + size].strip()
            if chunk:
                if len(chunks) >= max_chunks:
                    raise ValueError("DOCUMENT_CHUNK_LIMIT_EXCEEDED")
                chunks.append({"page": page, "text": chunk})
            if start + size >= len(text):
                break
    if not chunks:
        raise ValueError("EMPTY_DOCUMENT")
    return chunks, max(page_count, 1)
